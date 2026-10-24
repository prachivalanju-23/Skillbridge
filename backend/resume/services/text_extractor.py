import io

try:
    import pymupdf as fitz  # newer PyMuPDF
except ImportError:
    import fitz

from docx import Document


def extract_text(uploaded_file):
    name = uploaded_file.name.lower()
    data = uploaded_file.read()

    if name.endswith('.pdf'):
        pdf = fitz.open(stream=data, filetype='pdf')
        parts = []
        for page in pdf:
            # blocks sorted top to bottom, then left to right (fixes scrambled order)
            for block in page.get_text('blocks', sort=True):
                if block[6] == 0:  # text block
                    parts.append(block[4].strip())
        return '\n'.join(parts)

    if name.endswith('.docx'):
        doc = Document(io.BytesIO(data))
        parts = [p.text for p in doc.paragraphs]
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    parts.append(cell.text)
        return '\n'.join(parts)

    raise ValueError('Only PDF and DOCX files are supported')